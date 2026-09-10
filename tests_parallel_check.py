"""Headless tests for the multi-job engine. No network.

Covers the things that actually break under concurrency:
  1. two jobs writing the SAME sheet don't erase each other's statuses
  2. a job runs its batch queue back to back (finishing one starts the next)
  3. jobs on DIFFERENT sheets stay independent
  4. per-job model override reaches the client
  5. the 5-job ceiling is enforced
  6. the shared request budget caps total in-flight calls
  7. stopping one job leaves the others running

Run: .venv/bin/python tests_parallel_check.py
"""

import os
import queue
import shutil
import threading
import time
from pathlib import Path

import openpyxl

os.environ.setdefault("AI_API_KEY", "test-key")

import batch_runner
from batch_runner import JobManager, MAX_PARALLEL_JOBS, WorkbookRegistry, partition_batches

TMP = Path("/tmp/parallel_check")

CONFIG = {
    "ai_provider": {
        "base_url": "https://example.invalid/v1",
        "api_key_env": "AI_API_KEY",
        "model": "model-a:free",
        "temperature": 0.7,
        "max_retries": 0,
        "request_timeout_seconds": 5,
        "article_timeout_seconds": 60,
        "max_concurrent_requests": 2,
        "global_max_concurrent_requests": 3,
    },
    "free_models": ["model-a:free", "model-b:free"],
    "validation": {"word_count_tolerance_percent": 10, "max_correction_passes": 1},
    "abstract": {"min_words": 10, "max_words": 20},
    "agents": {},
    "writing_style": "ieee_paper",
    "max_consecutive_failures": 3,
    "excel_save_every_n_rows": 1,
}


def make_sheet(path: Path, rows: int):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Title", "Scope", "Status", "Notes"])
    for i in range(rows):
        ws.append([f"Article {i + 1}", f"Scope {i + 1}", None, None])
    wb.save(path)


def statuses_of(path: Path) -> dict:
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    hdr = [c.value for c in ws[1]]
    si = hdr.index("Status") + 1
    return {r: (ws.cell(row=r, column=si).value or "") for r in range(2, ws.max_row + 1)}


# --------------------------------------------------------------------- fakes

def fake_run_article_factory(delay=0.02, fail_rows=(), models_seen=None, inflight=None):
    """Replaces orchestrator.run_article so tests never touch the network."""
    def fake(row_number, title, scope, author, word_count, sections,
             config, client, output_dir, file_label=None):
        if models_seen is not None:
            models_seen.append(config["ai_provider"]["model"])
        if inflight is not None:
            inflight.enter()
        try:
            time.sleep(delay)
        finally:
            if inflight is not None:
                inflight.exit()

        from state import ArticleState
        state = ArticleState(row_number=row_number, title=title, scope=scope,
                             author=author, target_word_count=word_count,
                             target_sections=sections)
        if row_number in fail_rows:
            state.status = "Failed"
            state.notes = "synthetic failure"
        else:
            state.status = "Success"
            state.notes = f"Saved to {row_number}_x.pdf"
        return state
    return fake


class InflightTracker:
    def __init__(self):
        self.current = 0
        self.peak = 0
        self._lock = threading.Lock()

    def enter(self):
        with self._lock:
            self.current += 1
            self.peak = max(self.peak, self.current)

    def exit(self):
        with self._lock:
            self.current -= 1


def drain(manager, timeout=30):
    """Waits for all runners to finish."""
    deadline = time.time() + timeout
    while time.time() < deadline:
        if not manager.any_running():
            return True
        time.sleep(0.05)
    return False


# --------------------------------------------------------------------- tests

def test_same_sheet_concurrent_writes_not_lost():
    sheet = TMP / "shared.xlsx"
    make_sheet(sheet, 40)

    events = queue.Queue()
    batch_runner.run_article = fake_run_article_factory(delay=0.01)
    manager = JobManager(CONFIG, events.put)

    # Two jobs, same file, different batches (batch 1 = rows 2-11, batch 2 = rows 12-21).
    manager.add_job(sheet, TMP, "model-a:free", [1], 10, 500, 5)
    manager.add_job(sheet, TMP, "model-b:free", [2], 10, 500, 5)
    manager.start_all(CONFIG)
    assert drain(manager), "jobs did not finish"

    st = statuses_of(sheet)
    rows_batch1 = [r for r in range(2, 12) if st[r].startswith("Success")]
    rows_batch2 = [r for r in range(12, 22) if st[r].startswith("Success")]

    assert len(rows_batch1) == 10, f"batch 1 lost writes: {rows_batch1}"
    assert len(rows_batch2) == 10, f"batch 2 lost writes: {rows_batch2}"
    print(f"same-sheet parallel writes intact: batch1={len(rows_batch1)}/10 "
          f"batch2={len(rows_batch2)}/10: OK")


def test_batch_queue_runs_back_to_back():
    sheet = TMP / "queue.xlsx"
    make_sheet(sheet, 30)

    seen_batches = []
    events = queue.Queue()

    def emit(evt):
        if evt[0] == "job_batch":
            seen_batches.append(evt[2])
        events.put(evt)

    batch_runner.run_article = fake_run_article_factory(delay=0.005)
    manager = JobManager(CONFIG, emit)
    manager.add_job(sheet, TMP, "model-a:free", [1, 2, 3], 10, 500, 5)
    manager.start_all(CONFIG)
    assert drain(manager)

    assert seen_batches == [1, 2, 3], f"batches ran out of order/incomplete: {seen_batches}"
    st = statuses_of(sheet)
    done = sum(1 for v in st.values() if v.startswith("Success"))
    assert done == 30, f"expected all 30 rows done, got {done}"
    print(f"queued batches ran back to back {seen_batches}, all {done} rows done: OK")


def test_different_sheets_are_independent():
    a, b = TMP / "sheet_a.xlsx", TMP / "sheet_b.xlsx"
    make_sheet(a, 10)
    make_sheet(b, 10)

    batch_runner.run_article = fake_run_article_factory(delay=0.01)
    manager = JobManager(CONFIG, lambda e: None)
    manager.add_job(a, TMP, "model-a:free", [1], 10, 500, 5)
    manager.add_job(b, TMP, "model-b:free", [1], 10, 500, 5)
    manager.start_all(CONFIG)
    assert drain(manager)

    for sheet in (a, b):
        st = statuses_of(sheet)
        done = sum(1 for v in st.values() if v.startswith("Success"))
        assert done == 10, f"{sheet.name}: expected 10 done, got {done}"
    print("two different sheets processed independently: OK")


def test_per_job_model_override():
    sheet = TMP / "models.xlsx"
    make_sheet(sheet, 6)

    seen = []
    batch_runner.run_article = fake_run_article_factory(delay=0.005, models_seen=seen)
    manager = JobManager(CONFIG, lambda e: None)
    manager.add_job(sheet, TMP, "model-b:free", [1], 3, 500, 5)   # rows 2-4
    manager.add_job(sheet, TMP, "model-a:free", [2], 3, 500, 5)   # rows 5-7
    manager.start_all(CONFIG)
    assert drain(manager)

    assert "model-a:free" in seen and "model-b:free" in seen, seen
    print(f"per-job model override reached the pipeline ({sorted(set(seen))}): OK")


def test_max_five_jobs_enforced():
    sheet = TMP / "cap.xlsx"
    make_sheet(sheet, 5)
    manager = JobManager(CONFIG, lambda e: None)
    for _ in range(MAX_PARALLEL_JOBS):
        manager.add_job(sheet, TMP, "model-a:free", [1], 5, 500, 5)
    try:
        manager.add_job(sheet, TMP, "model-a:free", [1], 5, 500, 5)
    except ValueError as exc:
        assert "5" in str(exc), str(exc)
        print(f"job ceiling enforced at {MAX_PARALLEL_JOBS}: OK")
        return
    raise AssertionError("expected the 6th job to be rejected")


def test_global_request_budget_respected():
    """5 jobs must not exceed the configured global in-flight ceiling."""
    sheet = TMP / "budget.xlsx"
    make_sheet(sheet, 50)

    tracker = InflightTracker()
    batch_runner.run_article = fake_run_article_factory(delay=0.05, inflight=tracker)

    # The engine's semaphore guards AI calls; our fake stands in for one call per row,
    # so acquire it here the same way AIClient would.
    budget = CONFIG["ai_provider"]["global_max_concurrent_requests"]
    manager = JobManager(CONFIG, lambda e: None)

    real_fake = batch_runner.run_article

    def budgeted(*args, **kwargs):
        manager.request_semaphore.acquire()
        try:
            return real_fake(*args, **kwargs)
        finally:
            manager.request_semaphore.release()

    batch_runner.run_article = budgeted

    for i in range(MAX_PARALLEL_JOBS):
        manager.add_job(sheet, TMP, "model-a:free", [i + 1], 10, 500, 5)
    manager.start_all(CONFIG)
    assert drain(manager, timeout=60)

    assert tracker.peak <= budget, f"peak in-flight {tracker.peak} exceeded budget {budget}"
    print(f"global request budget held (peak {tracker.peak} <= {budget}) across "
          f"{MAX_PARALLEL_JOBS} jobs: OK")


def test_stopping_one_job_leaves_others_running():
    sheet_a, sheet_b = TMP / "stop_a.xlsx", TMP / "stop_b.xlsx"
    make_sheet(sheet_a, 30)
    make_sheet(sheet_b, 6)

    batch_runner.run_article = fake_run_article_factory(delay=0.08)
    manager = JobManager(CONFIG, lambda e: None)
    job_a = manager.add_job(sheet_a, TMP, "model-a:free", [1, 2, 3], 10, 500, 5)
    job_b = manager.add_job(sheet_b, TMP, "model-b:free", [1], 6, 500, 5)
    manager.start_all(CONFIG)

    time.sleep(0.2)
    manager.stop_job(job_a.job_id)
    assert drain(manager, timeout=30)

    assert job_a.status == "Stopped", job_a.status
    assert job_b.status == "Done", f"job B should have finished, got {job_b.status}"
    done_b = sum(1 for v in statuses_of(sheet_b).values() if v.startswith("Success"))
    assert done_b == 6, done_b
    print(f"stopping job A ({job_a.status}) left job B unaffected ({job_b.status}, "
          f"{done_b}/6 rows): OK")


def test_contention_tuning():
    """With several jobs running, each must take a fair share of the request budget and
    get a proportionally longer article deadline.

    Regression: 5 jobs × 2 section workers = 10 demands on a 6-slot budget. Requests
    queued on the semaphore and that wait counted against the article's wall clock, so
    articles 'timed out' against a perfectly healthy provider — one batch succeeded while
    five failed.
    """
    from batch_runner import JobRunner

    sheet = TMP / "contend.xlsx"
    make_sheet(sheet, 10)
    manager = JobManager(CONFIG, lambda e: None)
    job = manager.add_job(sheet, TMP, "model-a:free", [1], 10, 500, 5)

    budget = CONFIG["ai_provider"]["global_max_concurrent_requests"]   # 3 in tests
    base_timeout = CONFIG["ai_provider"]["article_timeout_seconds"]

    for parallel in (1, 2, 5):
        runner = JobRunner(job, CONFIG, manager.registry, lambda e: None,
                           manager.request_semaphore, running_jobs=lambda p=parallel: p)
        runner._tune_for_contention()
        workers = runner.config["ai_provider"]["max_concurrent_requests"]
        timeout = runner.config["ai_provider"]["article_timeout_seconds"]

        assert workers * parallel <= max(budget, parallel), \
            f"{parallel} jobs × {workers} workers oversubscribes the {budget}-slot budget"
        assert timeout >= base_timeout, (parallel, timeout)
        print(f"  {parallel} job(s): {workers} worker(s)/article, {timeout}s budget "
              f"(demand {workers * parallel} vs {budget} slots)")

    print("contention tuning keeps demand within the budget and extends deadlines: OK")


def test_already_done_rows_are_skipped():
    sheet = TMP / "skip.xlsx"
    make_sheet(sheet, 10)
    reg = WorkbookRegistry()
    for row in range(2, 7):
        reg.write_status(sheet, row, "Success", "pre-existing")

    ran = []
    base = fake_run_article_factory(delay=0.005)

    def tracking(row_number, *a, **kw):
        ran.append(row_number)
        return base(row_number, *a, **kw)

    batch_runner.run_article = tracking
    manager = JobManager(CONFIG, lambda e: None)
    manager.add_job(sheet, TMP, "model-a:free", [1], 10, 500, 5)
    manager.start_all(CONFIG)
    assert drain(manager)

    assert sorted(ran) == list(range(7, 12)), f"should only run unfinished rows, ran {sorted(ran)}"
    print(f"pre-finished rows skipped, ran only {sorted(ran)}: OK")


if __name__ == "__main__":
    if TMP.exists():
        shutil.rmtree(TMP)
    TMP.mkdir(parents=True)

    original = batch_runner.run_article
    try:
        test_same_sheet_concurrent_writes_not_lost()
        test_batch_queue_runs_back_to_back()
        test_different_sheets_are_independent()
        test_per_job_model_override()
        test_max_five_jobs_enforced()
        test_global_request_budget_respected()
        test_stopping_one_job_leaves_others_running()
        test_contention_tuning()
        test_already_done_rows_are_skipped()
    finally:
        batch_runner.run_article = original
        shutil.rmtree(TMP, ignore_errors=True)

    print("\nAll parallel-engine tests passed.")
