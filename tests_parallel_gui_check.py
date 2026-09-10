"""Headless checks for the reorganised UI:

  - batch-spec parsing
  - a sheet tab can launch several of ITS OWN batches in parallel
  - a different Excel sheet gets its own tab, sharing the app-wide job ceiling
  - parallel settings live on the Settings page and save correctly

Run: .venv/bin/python tests_parallel_gui_check.py
"""

import os
import shutil
import time
import tkinter as tk
from pathlib import Path
from tkinter import ttk

import openpyxl

os.environ.setdefault("AI_API_KEY", "test-key")

from dotenv import load_dotenv

load_dotenv()

import batch_runner
import gui_app
from batch_runner import MAX_PARALLEL_JOBS

TMP = Path("/tmp/parallel_gui_check")


def make_sheet(path: Path, rows: int):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Title", "Scope", "Status", "Notes"])
    for i in range(rows):
        ws.append([f"Article {i + 1}", f"Scope {i + 1}", None, None])
    wb.save(path)


def fake_run_article(row_number, title, scope, author, word_count, sections,
                     config, client, output_dir, file_label=None, year=None):
    from state import ArticleState
    time.sleep(0.01)
    st = ArticleState(row_number=row_number, title=title, scope=scope, author=author,
                      target_word_count=word_count, target_sections=sections)
    st.status = "Success"
    st.notes = f"Saved to {row_number}_x.pdf"
    return st


def test_batch_spec_parsing():
    cases = [
        ("1", 16, [1]), ("2,3", 16, [2, 3]), ("4-6", 16, [4, 5, 6]),
        ("1,3-5", 16, [1, 3, 4, 5]), ("all", 4, [1, 2, 3, 4]), ("", 3, [1, 2, 3]),
        ("2,2,2", 16, [2]), ("14-20", 16, [14, 15, 16]), ("99", 16, []), ("junk", 16, []),
    ]
    for spec, available, expected in cases:
        got = gui_app.parse_batch_spec(spec, available)
        assert got == expected, f"{spec!r}/{available} -> {got}, expected {expected}"
    print(f"batch spec parsing ({len(cases)} cases): OK")


def test_sheet_tab_runs_own_batches_in_parallel():
    sheet = TMP / "one_sheet.xlsx"
    make_sheet(sheet, 30)

    app = gui_app.App()
    app.update()

    tab = app.sheet_tabs[0]
    tab.excel_path = sheet
    tab.output_dir = TMP
    tab.batch_size_var.set("10")
    tab._load_rows()
    app.update()

    assert len(tab.batches) == 3, f"expected 3 batches of 10, got {len(tab.batches)}"

    tab.parallel_spec_var.set("1-3")
    tab.parallel_model_var.set("model-x:free")
    tab._start_selected_batches()
    app.update()

    assert len(app.job_manager.jobs) == 3, app.job_manager.jobs
    launched = sorted(j.batch_numbers[0] for j in app.job_manager.jobs.values())
    assert launched == [1, 2, 3], launched
    print(f"one sheet tab launched batches {launched} in parallel from its own tab: OK")

    deadline = time.time() + 30
    while time.time() < deadline and app.job_manager.any_running():
        app.update()
        time.sleep(0.05)

    wb = openpyxl.load_workbook(sheet)
    ws = wb.active
    hdr = [c.value for c in ws[1]]
    si = hdr.index("Status") + 1
    done = sum(1 for r in range(2, ws.max_row + 1)
               if str(ws.cell(row=r, column=si).value or "").startswith("Success"))
    assert done == 30, f"all 30 rows should be done across the 3 parallel batches, got {done}"
    print(f"all {done} rows completed across 3 concurrent batches, no lost writes: OK")

    app.destroy()


def test_second_sheet_gets_its_own_tab():
    sheet_a, sheet_b = TMP / "a.xlsx", TMP / "b.xlsx"
    make_sheet(sheet_a, 10)
    make_sheet(sheet_b, 10)

    app = gui_app.App()
    app.update()

    tab_a = app.sheet_tabs[0]
    tab_a.excel_path = sheet_a
    tab_a.output_dir = TMP
    tab_a._load_rows()
    app._rename_tab(tab_a, sheet_a.stem)

    app._add_sheet_tab()
    app.update()
    assert len(app.sheet_tabs) == 2, app.sheet_tabs

    tab_b = app.sheet_tabs[1]
    tab_b.excel_path = sheet_b
    tab_b.output_dir = TMP
    tab_b._load_rows()
    app._rename_tab(tab_b, sheet_b.stem)
    app.update()

    titles = [app.notebook.tab(i, "text") for i in range(app.notebook.index("end"))]
    assert titles[0] == "Settings", titles
    assert "a" in titles and "b" in titles, titles
    assert tab_a.job_manager is tab_b.job_manager, "sheet tabs must share one job manager"
    print(f"second sheet opened its own tab; tabs = {titles}: OK")

    # Ceiling is app-wide, not per tab.
    tab_a.parallel_spec_var.set("1")
    tab_a.parallel_model_var.set("m1:free")
    tab_a._start_selected_batches()
    tab_b.parallel_spec_var.set("1")
    tab_b.parallel_model_var.set("m2:free")
    tab_b._start_selected_batches()
    app.update()

    models = sorted(j.model for j in app.job_manager.jobs.values())
    assert models == ["m1:free", "m2:free"], models
    print(f"each sheet tab ran with its own model {models}: OK")

    deadline = time.time() + 30
    while time.time() < deadline and app.job_manager.any_running():
        app.update()
        time.sleep(0.05)
    app.destroy()


def test_parallel_settings_on_settings_page():
    app = gui_app.App()
    app.update()
    s = app.settings_tab

    for key in ("global_concurrency", "article_concurrency", "max_consecutive", "article_timeout"):
        assert key in s.vars, f"Settings page is missing the '{key}' field"
    print(f"parallel options present on Settings page: {sorted(k for k in s.vars if k in {'global_concurrency','article_concurrency','max_consecutive','article_timeout'})}: OK")

    original = config_manager_snapshot()
    try:
        s.vars["global_concurrency"].set("4")
        s.vars["max_consecutive"].set("2")
        s._save()
        saved = gui_app.config_manager.load_config()
        assert saved["ai_provider"]["global_max_concurrent_requests"] == 4, saved["ai_provider"]
        assert saved["max_consecutive_failures"] == 2, saved
        print("parallel settings persist to config.yaml: OK")
    finally:
        restore_config(original)
    app.destroy()


def config_manager_snapshot():
    return gui_app.config_manager.load_config()


def restore_config(cfg):
    gui_app.config_manager.save_config(cfg)


if __name__ == "__main__":
    if TMP.exists():
        shutil.rmtree(TMP)
    TMP.mkdir(parents=True)

    gui_app.messagebox.showwarning = lambda *a, **k: None
    gui_app.messagebox.showinfo = lambda *a, **k: None
    gui_app.messagebox.showerror = lambda *a, **k: None

    original_runner = batch_runner.run_article
    batch_runner.run_article = fake_run_article
    try:
        test_batch_spec_parsing()
        test_sheet_tab_runs_own_batches_in_parallel()
        test_second_sheet_gets_its_own_tab()
        test_parallel_settings_on_settings_page()
    finally:
        batch_runner.run_article = original_runner
        shutil.rmtree(TMP, ignore_errors=True)

    print("\nAll UI checks passed.")
